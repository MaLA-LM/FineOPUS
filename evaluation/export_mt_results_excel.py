#!/usr/bin/env python3
"""Export complete MT result JSON files to styled, auditable Excel workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from plot_mt_common import METRIC_COLUMNS, RESULTS_DIR, iteration_from_checkpoint


METRICS = list(METRIC_COLUMNS)
KEYS = ["model", "variant", "language", "checkpoint", "iteration"]
SUFFIX_LABELS = {
    "FineOPUS": "FineOPUS",
    "FineOPUS-High": "FineOPUS High",
    "FineOPUS-xHigh": "FineOPUS xHigh",
    "MaLA_Bi": "MaLA Bi",
    "MaLA_Bi_NLLB": "MaLA Bi + NLLB",
    "NLLB": "NLLB",
    **{f"FineOPUS-Stage{stage}": f"FineOPUS Stage {stage}" for stage in range(1, 5)},
    "FineOPUS-Stage4-High": "FineOPUS Stage 4 High",
    "FineOPUS-Stage4-xHigh": "FineOPUS Stage 4 xHigh",
}
DEFAULT_INPUTS = [
    RESULTS_DIR / "all_results_0.4B_multilingual_flores200.json",
    RESULTS_DIR / "all_results_0.4B_bilingual_flores200.json",
]


def variant_label(model: str) -> str:
    prefix = "0.4B_Pretrain_multilingual_"
    if model.startswith(prefix):
        suffix = model[len(prefix):]
    else:
        match = re.fullmatch(r"0.4B_Pretrain_eng_Latn-[a-z]{3}_[A-Za-z]+_(.+)", model)
        suffix = match.group(1) if match else model
    return SUFFIX_LABELS.get(suffix, suffix)


def direction_label(source: str, target: str) -> str:
    if source == "eng_Latn" and target != "eng_Latn":
        return "eng_to_x"
    if target == "eng_Latn" and source != "eng_Latn":
        return "x_to_eng"
    return "other"


def load_results(path: Path) -> tuple[list[dict], pd.DataFrame]:
    with path.open(encoding="utf-8") as handle:
        records = json.load(handle)
    if not isinstance(records, list) or not records or not all(isinstance(row, dict) for row in records):
        raise ValueError(f"Expected a nonempty list of result objects: {path}")
    raw = pd.DataFrame(records)
    required = {"model", "language", "checkpoint", "dataset", "source", "target", "status", *METRICS}
    if missing := required - set(raw.columns):
        raise ValueError(f"Missing required fields in {path}: {sorted(missing)}")
    if set(raw.columns) & {"variant", "iteration", "direction"}:
        raise ValueError("Input uses reserved derived fields: variant, iteration, direction")
    raw["variant"] = raw["model"].map(variant_label)
    raw["iteration"] = raw["checkpoint"].map(iteration_from_checkpoint)
    raw["direction"] = [direction_label(source, target) for source, target in zip(raw["source"], raw["target"])]
    leading = [
        "variant", "language", "checkpoint", "iteration", "dataset", "direction",
        "source", "target", *METRICS, "num_examples", "status", "model",
    ]
    columns = [column for column in leading if column in raw.columns]
    columns += [column for column in raw.columns if column not in columns]
    return records, raw[columns]


def rank_models(frame: pd.DataFrame, by_dataset: bool = False) -> pd.DataFrame:
    groups = ["dataset", "variant"] if by_dataset else ["variant"]
    ranking = frame.groupby(groups, sort=True, observed=True)[METRICS].mean()
    ranking["num_languages"] = frame.groupby(groups, observed=True)["language"].nunique()
    ranking = ranking.reset_index()
    rank_columns = []
    for metric in METRICS:
        column = metric + "_rank"
        values = ranking.groupby("dataset")[metric] if by_dataset else ranking[metric]
        ranking[column] = values.rank(ascending=False, method="min").astype(int)
        rank_columns.append(column)
    ranking["mean_metric_rank"] = ranking[rank_columns].mean(axis=1)
    values = ranking.groupby("dataset")["mean_metric_rank"] if by_dataset else ranking["mean_metric_rank"]
    ranking["overall_rank"] = values.rank(ascending=True, method="min").astype(int)
    order = (["dataset"] if by_dataset else []) + ["overall_rank", "variant"]
    return ranking.sort_values(order, kind="stable").reset_index(drop=True)


def summary_tables(raw: pd.DataFrame) -> dict[str, pd.DataFrame]:
    complete = raw.loc[raw["status"].eq("complete")].copy()
    if complete.empty:
        raise ValueError("No complete evaluations available for summary tables")
    if complete["direction"].eq("other").any():
        raise ValueError("Summary supports English-to-X and X-to-English evaluations only")
    if complete[KEYS + ["dataset", "direction"]].isna().any().any():
        raise ValueError("Complete records contain missing grouping fields")
    for metric in METRICS:
        complete[metric] = pd.to_numeric(complete[metric], errors="raise")
        if not complete[metric].map(lambda value: pd.notna(value) and math.isfinite(value)).all():
            raise ValueError(f"Complete records contain missing or non-finite {metric} scores")
    for field in ("bleu_tokenizer", "bleu_signature", "chrf_word_order", "comet_model", "few_shot", "limit"):
        if field in complete and complete[field].nunique(dropna=False) != 1:
            raise ValueError(f"Cannot average evaluations with different {field} settings")
    identity = ["model", "language", "checkpoint", "dataset", "source", "target"]
    if complete.duplicated(identity).any():
        raise ValueError("Duplicate complete evaluation directions; resolve them before aggregation")

    # Collapse dataset language aliases before equally weighting directions and datasets.
    cells = complete.groupby(KEYS + ["dataset", "direction"], sort=True)[METRICS].mean().reset_index()
    aliases = complete.groupby(KEYS + ["dataset", "direction"]).size().rename("num_alias_pairs").reset_index()
    cells = cells.merge(aliases, on=KEYS + ["dataset", "direction"], validate="one_to_one")
    dataset = cells.groupby(KEYS + ["dataset"], sort=True)[METRICS].mean().reset_index()
    direction_counts = cells.groupby(KEYS + ["dataset"])["direction"].nunique().rename("num_directions").reset_index()
    dataset = dataset.merge(direction_counts, on=KEYS + ["dataset"], validate="one_to_one")
    language = dataset.groupby(KEYS, sort=True)[METRICS].mean().reset_index()
    dataset_counts = dataset.groupby(KEYS)["dataset"].nunique().rename("num_datasets").reset_index()
    language = language.merge(dataset_counts, on=KEYS, validate="one_to_one")
    final = language.sort_values("iteration", kind="stable").groupby(["model", "language"], sort=False).tail(1)
    final = final.sort_values(["variant", "language"], kind="stable").reset_index(drop=True)
    final_dataset = dataset.merge(final[KEYS], on=KEYS, how="inner", validate="many_to_one")
    final_direction = cells.merge(final[KEYS], on=KEYS, how="inner", validate="many_to_one")

    coverage = raw.groupby(["variant", "model"], sort=True).agg(
        raw_rows=("model", "size"), num_checkpoints=("checkpoint", "nunique"),
        first_iteration=("iteration", "min"), last_iteration=("iteration", "max"),
        num_languages=("language", "nunique"), num_datasets=("dataset", "nunique"),
        num_directions=("direction", "nunique"),
    ).reset_index()
    status_counts = raw.assign(is_complete=raw["status"].eq("complete")).groupby("model")["is_complete"].sum()
    coverage["complete_rows"] = coverage["model"].map(status_counts)
    coverage["other_status_rows"] = coverage["raw_rows"] - coverage["complete_rows"]
    return {
        "RawData": raw,
        "Coverage": coverage,
        "LanguageByCheckpoint": language,
        "DatasetByCheckpoint": dataset,
        "FinalByLanguage": final,
        "FinalByDataset": final_dataset,
        "FinalByDirection": final_direction,
        "ModelRanking": rank_models(final),
        "DatasetRanking": rank_models(final_dataset, by_dataset=True),
    }


def overview(path: Path, raw: pd.DataFrame) -> pd.DataFrame:
    fields = [
        ("Source JSON", str(path.resolve())),
        ("Source SHA256", hashlib.sha256(path.read_bytes()).hexdigest()),
        ("Exported UTC", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("Raw records", len(raw)),
        ("Complete records", int(raw["status"].eq("complete").sum())),
        ("Actual model names", raw["model"].nunique()),
        ("Model variants", raw["variant"].nunique()),
        ("Languages", raw["language"].nunique()),
        ("Datasets", ", ".join(sorted(raw["dataset"].unique()))),
        ("RawData", "All input records and original fields, in source order; variant, iteration and direction are derived columns. No plotting exclusions are applied."),
        ("Coverage", "Raw/complete record counts and observed checkpoint, language, dataset and direction counts per actual model."),
        ("LanguageByCheckpoint", "Each model/language/checkpoint: mean of available dataset scores. num_datasets records actual coverage."),
        ("DatasetByCheckpoint", "Each model/language/checkpoint/dataset: equal-weight mean of available English-to-X and X-to-English directions. num_directions records actual coverage."),
        ("FinalByLanguage", "Latest complete checkpoint observed per actual model and language, not the highest-scoring checkpoint or a verified end-of-training marker."),
        ("FinalByDataset", "Dataset scores at the same checkpoint selected in FinalByLanguage; no independently selected checkpoint per dataset."),
        ("FinalByDirection", "Direction scores at the selected final checkpoint, with dataset language aliases averaged first. num_alias_pairs shows the number of source rows averaged."),
        ("ModelRanking", "Equal-weight mean of FinalByLanguage across available languages per variant. Higher scores are better; each metric is ranked separately, then overall_rank ranks the mean of the three metric ranks (ties use minimum rank)."),
        ("DatasetRanking", "Same ranking method independently per dataset, using equal-weight language means. num_languages records coverage."),
        ("Aggregation order", "Dataset aliases -> directions -> available datasets -> languages; no weighting by sentence count. Bilingual Arabic BOUQuET varieties are averaged within each direction first."),
        ("Coverage caution", "Missing datasets, directions or languages are not assigned zero. Rankings use observed coverage, not an enforced common-language subset; compare num_languages and num_datasets."),
        ("Metric columns", "bleu = spBLEU-200 when bleu_tokenizer is flores200; chrf = chrF++ when chrf_word_order is 2; comet remains on its original scale (not multiplied by 100)."),
        ("Precision and blanks", "Number formats affect display only; scores are not explicitly rounded. JSON null and empty strings are blank in Excel; the original JSON remains authoritative for those distinctions."),
        ("Regeneration", "This workbook is a static snapshot. Rerun evaluation/export_mt_results_excel.py after the source JSON changes."),
    ]
    return pd.DataFrame(fields, columns=["Field", "Description"])


def column_width(column: str) -> int:
    widths = {
        "variant": 27, "language": 16, "checkpoint": 19, "dataset": 24,
        "direction": 16, "source": 16, "target": 16, "model": 64,
        "result_dir": 86, "reason": 48, "bleu_signature": 68, "comet_model": 31,
        "bleu_tokenizer": 19, "status": 14, "Field": 27, "Description": 112,
    }
    return widths.get(column, min(27, max(15, len(column) + 3)))


def style_sheet(sheet, columns: list[str], index: int) -> None:
    sheet.freeze_panes = "D2" if sheet.title == "RawData" else "A2"
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.row_dimensions[1].height = 30
    sheet.sheet_properties.tabColor = "24756A" if sheet.title.startswith("Final") else "305C80"
    for number, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(number)].width = column_width(column)
        cell = sheet.cell(1, number)
        cell.fill = PatternFill("solid", fgColor="305C80")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell, column in zip(row, columns):
            # External result strings must remain text, even when beginning with '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=sheet.title == "Overview")
            if column == "comet":
                cell.number_format = "0.000000"
            elif column in ("bleu", "chrf", "seconds", "comet_seconds", "mean_metric_rank"):
                cell.number_format = "0.000"
            elif column.endswith("_rank") or column.startswith("num_") or column in (
                "iteration", "first_iteration", "last_iteration", "raw_rows", "complete_rows",
                "other_status_rows", "few_shot", "limit", "chrf_word_order",
            ):
                cell.number_format = "0"
    if sheet.title == "Overview":
        for row in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 48 if row >= 11 else 30
    if sheet.max_row > 1:
        table = Table(displayName=f"T{index:02d}_{sheet.title}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        if sheet.title.startswith("Final") or sheet.title.endswith("Ranking"):
            for metric in METRICS:
                if metric in columns:
                    letter = get_column_letter(columns.index(metric) + 1)
                    sheet.conditional_formatting.add(
                        f"{letter}2:{letter}{sheet.max_row}",
                        ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="9FD8BE"),
                    )


def export_workbook(source: Path, output: Path) -> dict[str, int]:
    records, raw = load_results(source)
    sheets = {"Overview": overview(source, raw), **summary_tables(raw)}
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = tempfile.NamedTemporaryFile(prefix=output.stem + "_", suffix=".tmp.xlsx", dir=output.parent, delete=False)
    temporary_path = Path(temporary.name)
    temporary.close()
    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            writer.book.properties.title = source.stem
            writer.book.properties.subject = "MT evaluation: full results, checkpoints and language-macro rankings"
            writer.book.properties.creator = "FineOPUS"
            for index, (name, frame) in enumerate(sheets.items(), 1):
                if len(frame) > 1_048_575 or len(frame.columns) > 16_384:
                    raise ValueError(f"Sheet exceeds Excel dimensions: {name}")
                frame.to_excel(writer, sheet_name=name, index=False)
                style_sheet(writer.sheets[name], list(frame.columns), index)
        temporary_path.replace(output)
    finally:
        temporary_path.unlink(missing_ok=True)
    counts = {name: len(frame) for name, frame in sheets.items()}
    assert counts["RawData"] == len(records)
    print(f"Wrote {output}: {counts}", flush=True)
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="*", type=Path, help="JSON files (default: both 0.4B flores200 aggregates)")
    parser.add_argument("--output-dir", type=Path, help="Destination directory; default: beside each source JSON")
    args = parser.parse_args()
    for source in args.inputs or DEFAULT_INPUTS:
        source = source.resolve()
        output = (args.output_dir / source.with_suffix(".xlsx").name) if args.output_dir else source.with_suffix(".xlsx")
        export_workbook(source, output)


if __name__ == "__main__":
    main()
